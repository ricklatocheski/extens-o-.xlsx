#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_demo.py
=============
Gera uma pasta de trabalho de demonstracao (dist/DemoTemaEscuro.xlsx) com o
tema escuro ja aplicado ao conteudo, usando as mesmas cores do
DarkTheme.thmx. Serve para visualizar o resultado do suplemento sem precisar
executar as macros.

Uso:
    python3 build_demo.py     (requer: pip install openpyxl)
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DIST_DIR = os.path.join(BASE_DIR, "dist")
OUTPUT = os.path.join(DIST_DIR, "DemoTemaEscuro.xlsx")

# Paleta (identica ao tema .thmx e ao suplemento VBA)
FUNDO = "1E1E1E"
FUNDO_ALT = "2D2D30"
CABECALHO = "2A2A2A"
TEXTO = "F2F2F2"
ACENTO = "4FC3F7"
BORDA = "3F3F3F"


def main():
    os.makedirs(DIST_DIR, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Demonstracao"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ACENTO

    fundo = PatternFill("solid", fgColor=FUNDO)
    fundo_alt = PatternFill("solid", fgColor=FUNDO_ALT)
    cab = PatternFill("solid", fgColor=CABECALHO)
    borda = Border(*[Side(style="thin", color=BORDA)] * 4)

    dados = [
        ("Produto", "Categoria", "Preco", "Estoque"),
        ("Teclado", "Perifericos", 149.90, 32),
        ("Monitor", "Telas", 899.00, 12),
        ("Mouse", "Perifericos", 79.90, 54),
        ("Notebook", "Computadores", 4299.00, 7),
        ("Webcam", "Perifericos", 199.90, 21),
    ]

    for r, linha in enumerate(dados, start=1):
        for c, valor in enumerate(linha, start=1):
            cel = ws.cell(row=r, column=c, value=valor)
            cel.border = borda
            cel.alignment = Alignment(vertical="center")
            if r == 1:
                cel.fill = cab
                cel.font = Font(name="Calibri", color=ACENTO, bold=True, size=12)
            else:
                cel.fill = fundo if r % 2 == 0 else fundo_alt
                cel.font = Font(name="Calibri", color=TEXTO)

    # Preenche uma margem extra de celulas para dar o efeito de "modo escuro"
    for r in range(1, 30):
        for c in range(1, 8):
            cel = ws.cell(row=r, column=c)
            if cel.fill.fgColor.rgb in (None, "00000000"):
                cel.fill = fundo

    widths = {"A": 16, "B": 16, "C": 12, "D": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    wb.save(OUTPUT)
    print(f"[OK] Demonstracao gerada: {OUTPUT}")


if __name__ == "__main__":
    main()
